# -*- coding: utf-8 -*-
"""Orquestra a automação de ponta a ponta."""
from __future__ import annotations

from dataclasses import dataclass, field
from pathlib import Path

import openpyxl

from .compositions import CompositionIndex, build_index
from .config import Config
from .layout import SheetLayout, TemplateLayout, detect, detect_template, is_empty
from .ooxml import Workbook, ensure_row, retarget_sumproduct, set_cell, show_row
from .pdf_budget import Topic, items_by_code, read_budget
from .planner import SheetPlan, plan_sheet
from .report import Audit, build_rows, log_sheet_xml, text_report
from .synth import SynthPlan, plan_synthesis, write_synthesis
from .resolver import Resolver
from .textutil import normalize, sanitize_sheet_name, similarity


@dataclass
class Result:
    output: Path
    plans: list[SheetPlan] = field(default_factory=list)
    synth: list[SynthPlan] = field(default_factory=list)
    topics: list[Topic] = field(default_factory=list)
    index: CompositionIndex | None = None
    audit: Audit | None = None

    @property
    def report(self) -> str:
        return text_report(self.audit) if self.audit else ""


MIN_TOKEN = 5


def topic_sheet_score(topic_name: str, sheet_name: str) -> float:
    """Quão bem o nome de um tópico do orçamento casa com o de uma aba.

    Comparar por substring seria frágil: a aba "RES" casaria com o tópico
    "SERVIÇOS PRELIMINARES". O casamento é por token, aceitando prefixo, que é
    o que liga "REVESTIMENTO" a "REVESTIMENTOS" e "PAREDES E PAINÉIS" a
    "PAREDES" sem aceitar coincidências curtas.
    """
    topic_tokens = [t for t in normalize(topic_name).split() if len(t) >= MIN_TOKEN]
    sheet_tokens = [t for t in normalize(sheet_name).split() if len(t) >= MIN_TOKEN]
    base = similarity(topic_name, sheet_name)
    if not topic_tokens or not sheet_tokens:
        return base

    def matches(a: str, b: str) -> bool:
        return a.startswith(b) or b.startswith(a)

    covered = sum(any(matches(s, t) for t in topic_tokens) for s in sheet_tokens)
    return max(base, 0.95) if covered == len(sheet_tokens) else base


def match_topics(topics: list[Topic], sheet_names: list[str], config: Config,
                 usable=None) -> list[tuple[Topic, str]]:
    """Casa cada tópico do orçamento com a aba de destino correspondente.

    Primeiro o mapa explícito da configuração; depois semelhança de nome.
    `usable(nome)` permite descartar abas sem layout de levantamento.
    """
    wanted = config.targets.sheets
    manual = {str(k): v for k, v in config.targets.topic_map.items()}
    taken: set[str] = set()
    pairs: list[tuple[Topic, str]] = []
    for topic in topics:
        target = manual.get(str(topic.number))
        explicit = target is not None
        if target is None:
            best, best_score = None, 0.0
            for name in sheet_names:
                if (name in taken or (wanted and name not in wanted)
                        or config.targets.suffix.strip() in name):
                    continue
                if usable is not None and not usable(name):
                    continue
                score = topic_sheet_score(topic.name, name)
                if score > best_score:
                    best, best_score = name, score
            target = best if best_score >= config.targets.min_topic_similarity else None
        if target is None:
            continue
        if wanted and target not in wanted:
            continue
        if target not in sheet_names:
            raise ValueError(f"aba de destino inexistente: {target!r}")
        if not explicit and usable is not None and not usable(target):
            continue
        taken.add(target)
        pairs.append((topic, target))
    return pairs


def run(xlsx_path: str | Path, pdf_path: str | Path | None, output: str | Path,
        config: Config | None = None, topics: list[Topic] | None = None) -> Result:
    """Gera o arquivo com as abas (AUTO).

    `topics` permite injetar um orçamento já interpretado, dispensando o PDF —
    é o que torna a automação testável de ponta a ponta sem depender do
    pdfplumber nem de um arquivo binário de apoio.
    """
    config = config or Config()
    xlsx_path, output = Path(xlsx_path), Path(output)
    if topics is None:
        if pdf_path is None:
            raise ValueError("informe `pdf_path` ou `topics`")
        topics = read_budget(Path(pdf_path), config.pdf)
    budget_codes = set(items_by_code(topics))

    # `read_only` é o que torna isto viável em servidor: numa pasta com uma aba
    # de 65 mil linhas, a carga normal do openpyxl materializa tudo e passa de
    # 450 MB por workbook — e são dois. Em streaming, o par fica em torno de
    # 50 MB. A detecção de layout já lê por varredura única e limitada, então
    # não depende de acesso aleatório a célula, que em read_only seria O(n²).
    formulas = openpyxl.load_workbook(xlsx_path, read_only=True)
    values = openpyxl.load_workbook(xlsx_path, data_only=True, read_only=True)
    index = build_index(values, config.compositions)
    resolver = Resolver(index, config.compositions)

    def usable(name: str) -> bool:
        try:
            detect(formulas[name], values[name], config.targets)
            return True
        except Exception:
            return False

    pairs = match_topics(topics, list(formulas.sheetnames), config, usable)
    template = config.targets.template_sheet
    can_generate = bool(config.targets.generate_missing and template
                        and template in formulas.sheetnames)
    if not pairs and not can_generate:
        # sem aba de destino e sem molde não há o que fazer: ou o arquivo não é
        # um levantamento, ou o nome do molde está errado na configuração
        raise ValueError(
            "nenhum tópico do orçamento casou com uma aba do arquivo e não há aba-molde "
            f"{template!r} para sintetizar; declare `targets.topic_map` ou "
            "`targets.template_sheet` na configuração")

    package = Workbook(xlsx_path)
    plans: list[SheetPlan] = []
    for topic, sheet_name in pairs:
        layout = detect(formulas[sheet_name], values[sheet_name], config.targets)
        plan = plan_sheet(layout, topic.items, index, resolver, budget_codes,
                          topic.number, topic.name, config.rules)
        plans.append(plan)
        _write_sheet(package, plan, layout, config)

    # ---- tópicos sem aba de destino: a aba é sintetizada a partir do molde
    synth = _synthesize_missing(package, formulas, values, topics, pairs,
                                index, resolver, config)

    audit = Audit(plans, config.targets.suffix, budget_codes, synth)
    if config.rules.write_log_sheet:
        bold = package.append_style(bold=True)
        wrap = package.append_style(wrap=True)
        package.add_sheet(config.rules.log_sheet_name,
                          log_sheet_xml(build_rows(audit), bold, wrap))

    package.save(output)
    formulas.close()
    values.close()
    return Result(output=output, plans=plans, synth=synth, topics=topics,
                  index=index, audit=audit)


def _synthesize_missing(package: Workbook, formulas, values, topics: list[Topic],
                        pairs, index, resolver, config: Config) -> list[SynthPlan]:
    """Cria uma aba para cada tópico do orçamento que não tem aba de destino.

    Sem molde declarado, ou com `generate_missing` desligado, o modo novo
    simplesmente não roda e o comportamento anterior é preservado por inteiro.
    """
    template = config.targets.template_sheet
    if not (config.targets.generate_missing and template and template in formulas.sheetnames):
        return []
    layout = detect_template(formulas[template], values[template], config.targets)

    handled = {topic.number for topic, _ in pairs}
    taken = set(formulas.sheetnames) | {sheet.name for sheet in package.pending()}
    produced: list[SynthPlan] = []
    for topic in topics:
        if topic.number in handled:
            continue
        name = sanitize_sheet_name(topic.name, taken)
        plan = plan_synthesis(layout, topic, name, index, resolver, config)
        if not plan.rows or not plan.columns:
            # tópico sem nada a levantar (nenhum item com composição, ou nenhum
            # insumo nas seções configuradas): registra e não cria aba vazia
            plan.created = False
            plan.reason = ("nenhum item com composição no arquivo" if not plan.rows
                           else "nenhum insumo nas seções configuradas em "
                                "`compositions.column_sections`")
            produced.append(plan)
            continue
        taken.add(name)
        sheet = package.clone_sheet(template, name)
        sheet.xml = write_synthesis(sheet.xml, plan, layout, config)
        produced.append(plan)
    return produced


def _style_template(layout: SheetLayout) -> int | None:
    """Linha existente que serve de molde visual para linhas criadas do zero."""
    for row in reversed(layout.rows_inside_total()):
        if not layout.rows[row].is_blank:
            return row
    return layout.first_row


def _write_sheet(package: Workbook, plan: SheetPlan, layout: SheetLayout, config: Config) -> None:
    sheet = package.clone_sheet(plan.sheet, plan.sheet + config.targets.suffix)
    xml = sheet.xml
    template = _style_template(layout)
    for planned in plan.ordered():
        row = planned.row
        xml = ensure_row(xml, row, template)
        if planned.writes_identity:
            if layout.rows[row].is_blank:
                xml = show_row(xml, row)
            else:
                # a linha passa a ser outro serviço: limpa os insumos rastreados que
                # a nova composição não possui, para não sobrar valor do código antigo
                for column in layout.tracked:
                    if column not in planned.coefficients:
                        xml = set_cell(xml, row, column)
            xml = set_cell(xml, row, layout.order_column, text=planned.item.order)
            xml = set_cell(xml, row, layout.code_column, text=planned.item.code)
            xml = set_cell(xml, row, layout.description_column,
                           text=planned.composition.description)
            xml = set_cell(xml, row, layout.unit_column,
                           text=planned.composition.unit or planned.item.unit or "")
            if planned.quantity is not None:
                xml = set_cell(xml, row, layout.quantity_column, number=planned.quantity)
        else:
            if planned.quantity_source == "orcamento" and planned.quantity is not None:
                xml = set_cell(xml, row, layout.quantity_column, number=planned.quantity)
            if config.rules.fill_empty_identity:
                # a linha já traz o código certo, mas a planilha original pode ter
                # a unidade/descrição vazia ou presa a um VLOOKUP quebrado (#REF!).
                # Preenche apenas nesse caso — conteúdo válido nunca é tocado.
                existing = layout.rows[row]
                if existing.unit is None and planned.composition.unit:
                    xml = set_cell(xml, row, layout.unit_column,
                                   text=planned.composition.unit)
                    planned.notes.append("unidade preenchida (estava vazia na aba)")
                if is_empty(existing.description) and planned.composition.description:
                    xml = set_cell(xml, row, layout.description_column,
                                   text=planned.composition.description)
                    planned.notes.append("descrição preenchida (estava vazia na aba)")
        for column, coefficient in planned.coefficients.items():
            xml = set_cell(xml, row, column,
                           formula=coefficient.formula(config.compositions.coefficient_column),
                           value=coefficient.value)
    if plan.new_last_row:
        xml = retarget_sumproduct(xml, layout.total_row, layout.last_row,
                                  plan.new_last_row, list(layout.tracked))
    sheet.xml = xml
