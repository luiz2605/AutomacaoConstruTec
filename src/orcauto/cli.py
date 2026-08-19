# -*- coding: utf-8 -*-
"""Linha de comando do orcauto."""
from __future__ import annotations

import argparse
import sys
import zipfile
from pathlib import Path

from .config import Config


def _add_common(parser: argparse.ArgumentParser) -> None:
    parser.add_argument("--xlsx", required=True, help="planilha de levantamento (.xlsx)")
    parser.add_argument("--config", help="arquivo .toml/.json de configuração")


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(
        prog="orcauto",
        description="Cruza um orçamento analítico em PDF com as composições de custo de uma "
                    "planilha Excel e gera versões automatizadas das abas de levantamento.")
    sub = parser.add_subparsers(dest="command", required=True)

    run = sub.add_parser("run", help="gera o arquivo com as abas (AUTO)")
    _add_common(run)
    run.add_argument("--pdf", required=True, help="orçamento analítico (.pdf)")
    run.add_argument("--out", required=True, help="arquivo de saída (.xlsx)")
    run.add_argument("--quiet", action="store_true", help="não imprime o relatório")

    inspect = sub.add_parser("inspect", help="mostra o que foi detectado, sem gravar nada")
    _add_common(inspect)
    inspect.add_argument("--pdf", help="opcional: também lista os tópicos do orçamento")

    check = sub.add_parser("check", help="confere um arquivo gerado contra o original")
    check.add_argument("--original", required=True)
    check.add_argument("--generated", required=True)
    return parser


def command_run(args) -> int:
    from .pipeline import run
    result = run(args.xlsx, args.pdf, args.out, Config.load(args.config))
    if not args.quiet:
        print(result.report)
    criadas = [plan for plan in result.synth if plan.created]
    written = (sum(len(plan.rows) for plan in result.plans)
               + sum(len(plan.rows) for plan in criadas))
    skipped = (sum(len(plan.skipped) for plan in result.plans)
               + sum(len(plan.skipped) for plan in result.synth))
    print(f"gerado: {result.output}")
    print(f"  {len(result.plans)} aba(s) preenchida(s), {len(criadas)} aba(s) criada(s) do molde")
    print(f"  {written} linhas gravadas, {skipped} itens não aplicados")
    nao_criadas = [plan for plan in result.synth if not plan.created]
    for plan in nao_criadas:
        print(f"  aba não criada — tópico {plan.topic_number} {plan.topic_name}: {plan.reason}")
    return 0


def command_inspect(args) -> int:
    import openpyxl

    from .compositions import autodetect_sheets, build_index
    from .layout import detect, detect_template
    from .pipeline import match_topics

    config = Config.load(args.config)
    formulas = openpyxl.load_workbook(args.xlsx)
    values = openpyxl.load_workbook(args.xlsx, data_only=True)

    detected = config.compositions.sheets or autodetect_sheets(values, config.compositions)
    print("abas de composição:", ", ".join(detected) or "(nenhuma)")
    index = build_index(values, config.compositions, detected)
    print(f"composições indexadas: {len(index)}\n")

    usable = []
    for name in formulas.sheetnames:
        try:
            layout = detect(formulas[name], values[name], config.targets)
        except Exception:
            continue
        usable.append(name)
        print(f"[{name}]")
        print(f"   cabeçalho linha {layout.header_row} | insumos linha {layout.input_row} | "
              f"quantidade coluna {layout.quantity_column}")
        print(f"   bloco de serviços {layout.first_row}-{layout.last_row} | "
              f"TOTAL linha {layout.total_row} | livres {layout.free_rows()}")
        print(f"   insumos rastreados: {layout.tracked}")
    if not usable:
        print("nenhuma aba com layout de levantamento pronto "
              "(normal em arquivo-base: as abas serão criadas do molde)")

    template = config.targets.template_sheet
    molde = None
    if config.targets.generate_missing and template in formulas.sheetnames:
        try:
            molde = detect_template(formulas[template], values[template], config.targets)
        except Exception as error:
            print(f"\naba-molde {template!r}: NÃO utilizável — {error}")
    elif template not in formulas.sheetnames:
        print(f"\naba-molde {template!r} não existe no arquivo "
              "(só o modo de preenchimento vai rodar)")
    if molde is not None:
        print(f"\naba-molde [{template}]")
        print(f"   cabeçalho linha {molde.header_row} | insumos linha {molde.input_row} | "
              f"quantidade coluna {molde.quantity_column}")
        print(f"   bloco de serviços {molde.first_row}-{molde.last_row} | "
              f"TOTAL linha {molde.total_row}")
        print(f"   colunas já formatadas: {', '.join(molde.slots)} "
              f"({len(molde.slots)}; o que passar disso é acrescentado à direita)")

    if args.pdf:
        from .pdf_budget import read_budget
        topics = read_budget(args.pdf, config.pdf)
        print(f"\ntópicos do orçamento: {len(topics)}")
        pairs = dict((topic.number, sheet) for topic, sheet in
                     match_topics(topics, usable, config,
                                  lambda name: name in usable))
        from .textutil import sanitize_sheet_name
        taken = set(formulas.sheetnames)
        for topic in topics:
            destino = pairs.get(topic.number)
            if destino:
                destino = f"->  {destino} (preencher)"
            elif molde is not None:
                nome = sanitize_sheet_name(topic.name, taken)
                taken.add(nome)
                destino = f"->  {nome} (criar do molde)"
            else:
                destino = "sem destino"
            print(f"   {topic.number:>3} {topic.name:<34} {len(topic.items):>3} itens  {destino}")
    return 0


def command_check(args) -> int:
    """Confere que as abas originais não foram tocadas e que o pacote está íntegro."""
    from xml.etree import ElementTree

    original, generated = zipfile.ZipFile(args.original), zipfile.ZipFile(args.generated)
    names_o, names_g = set(original.namelist()), set(generated.namelist())
    allowed = {"[Content_Types].xml", "xl/_rels/workbook.xml.rels",
               "xl/workbook.xml", "xl/styles.xml"}
    changed = {n for n in names_o & names_g if original.read(n) != generated.read(n)}
    unexpected = changed - allowed
    malformed = []
    for name in names_g:
        if name.endswith((".xml", ".rels")):
            try:
                ElementTree.fromstring(generated.read(name))
            except ElementTree.ParseError as error:
                malformed.append((name, str(error)))

    print(f"partes novas ......... {len(names_g - names_o)}")
    print(f"partes removidas ..... {sorted(names_o - names_g) or 'nenhuma'}")
    print(f"originais alteradas .. {sorted(changed) or 'nenhuma'}")
    print(f"XML mal formado ...... {malformed or 'nenhum'}")
    ok = not unexpected and not malformed
    if unexpected:
        print(f"\nFALHOU: partes originais alteradas fora do esperado: {sorted(unexpected)}")
    print("\nOK: abas originais preservadas e pacote íntegro." if ok else "\nFALHOU.")
    return 0 if ok else 1


def main(argv: list[str] | None = None) -> int:
    args = build_parser().parse_args(argv)
    handlers = {"run": command_run, "inspect": command_inspect, "check": command_check}
    try:
        return handlers[args.command](args)
    except Exception as error:                      # mensagem limpa, sem traceback
        print(f"erro: {error}", file=sys.stderr)
        return 2


if __name__ == "__main__":                          # pragma: no cover
    raise SystemExit(main())
