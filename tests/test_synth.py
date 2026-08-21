"""Fase 2 — síntese de aba a partir do molde."""
import openpyxl
import pytest

from orcauto.compositions import CompositionIndex, parse_sheet
from orcauto.config import CompositionConfig, Config, TargetConfig
from orcauto.layout import LayoutError, TemplateLayout, detect, detect_template
from orcauto.ooxml import total_formula, unit_lookup_formula
from orcauto.resolver import Resolver, topic_inputs
from orcauto.synth import plan_synthesis, write_synthesis
from orcauto.textutil import sanitize_sheet_name


# ----------------------------------------------------------------- D6: nomes
def test_sanitiza_caracteres_proibidos_pelo_excel():
    assert sanitize_sheet_name("ABC/DEF:GHI[JKL]") == "ABC DEF GHI JKL"
    assert sanitize_sheet_name("a\\b?c*d") == "a b c d"


def test_trunca_em_31_na_fronteira_de_palavra():
    nome = sanitize_sheet_name("SERVICOS DE INSTALACAO HIDRAULICA COMPLETA")
    assert len(nome) <= 31
    assert not nome.endswith(" ")
    assert nome == "SERVICOS DE INSTALACAO"          # não corta no meio da palavra


def test_trunca_duro_quando_nao_ha_espaco_util():
    nome = sanitize_sheet_name("A" * 60)
    assert nome == "A" * 31


def test_nome_vazio_vira_rotulo_generico():
    assert sanitize_sheet_name("   ") == "ABA"
    assert sanitize_sheet_name(":::") == "ABA"


def test_resolve_colisao_com_sufixo_numerico():
    assert sanitize_sheet_name("PAREDES", {"PAREDES"}) == "PAREDES 2"
    assert sanitize_sheet_name("PAREDES", {"PAREDES", "PAREDES 2"}) == "PAREDES 3"


def test_colisao_respeita_o_limite_de_31():
    longo = "INSTALACOES HIDROSSANITARIAS X"
    nome = sanitize_sheet_name(longo, {longo})
    assert len(nome) <= 31 and nome.endswith(" 2")


# ------------------------------------------------- fórmula de TOTAL gerada
def test_total_formula_sumproduct_simples():
    assert total_formula("F", "E", 10, 13) == "SUMPRODUCT($E$10:$E$13,F10:F13)"


def test_total_formula_com_arredondamento_de_saco():
    assert total_formula("G", "E", 10, 13, bag_size=50) == \
        "ROUNDUP((SUMPRODUCT($E$10:$E$13,G10:G13)/50),0)"


def test_unit_lookup_formula_segue_o_padrao_do_molde():
    assert unit_lookup_formula("H", 7, "insumos!$A$6:$D$8635") == \
        'IFERROR(VLOOKUP(H7,insumos!$A$6:$D$8635,3,FALSE),"")'


# --------------------------------------------------- colunas a partir do tópico
@pytest.fixture
def resolver(composition_rows):
    index = CompositionIndex(parse_sheet(composition_rows, "COMPOSICOES"))
    return index, Resolver(index, CompositionConfig(service_code_re=r"^S"))


def test_colunas_vem_da_resolucao_recursiva_nao_do_primeiro_nivel(resolver, topics):
    """S002 consome o serviço S003; quem vira coluna são os insumos DE DENTRO.

    É o caso real do C0329, que traz C3129 no bloco SERVIÇOS. Lendo só o
    primeiro nível, o código do sub-serviço viraria uma coluna e os materiais
    de verdade sumiriam.
    """
    index, resolve = resolver
    colunas = topic_inputs(topics[0].items, index, resolve)
    codigos = [c.code for c in colunas]
    assert "S003" not in codigos                  # sub-composição não vira coluna
    assert "X001" in codigos and "X002" in codigos


def test_ordem_das_colunas_e_deterministica(resolver, topics):
    index, resolve = resolver
    primeira = [c.code for c in topic_inputs(topics[0].items, index, resolve)]
    segunda = [c.code for c in topic_inputs(topics[0].items, index, resolve)]
    assert primeira == segunda
    assert primeira[0] == "X001"                  # ordem de primeira aparição


def test_coluna_guarda_descricao_e_unidade(resolver, topics):
    index, resolve = resolver
    coluna = {c.code: c for c in topic_inputs(topics[0].items, index, resolve)}["X001"]
    assert coluna.description == "CIMENTO" and coluna.unit == "KG"
    assert "S001" in coluna.used_by


# ------------------------------------------------------- detecção do molde
def test_detect_template_le_as_ancoras_do_molde(template_workbook_path):
    formulas = openpyxl.load_workbook(template_workbook_path)
    values = openpyxl.load_workbook(template_workbook_path, data_only=True)
    with pytest.raises(LayoutError, match="nenhum insumo rastreado"):
        detect(formulas["MODELO BASE"], values["MODELO BASE"])
    layout = detect_template(formulas["MODELO BASE"], values["MODELO BASE"])
    assert (layout.header_row, layout.input_row) == (8, 7)
    assert (layout.first_row, layout.last_row, layout.total_row) == (10, 12, 14)
    assert layout.slots == ["F", "G"]              # deduzidas do SUMPRODUCT do TOTAL


# --------------------------------------------------------- ponta a ponta
@pytest.fixture
def gerado(template_workbook_path, topics, tmp_path):
    from orcauto.pipeline import run
    config = Config(compositions=CompositionConfig(sheets=["COMPOSICOES"],
                                                   service_code_re=r"^S"),
                    targets=TargetConfig(template_sheet="MODELO BASE"))
    destino = tmp_path / "gerado.xlsx"
    resultado = run(template_workbook_path, None, destino, config, topics=topics)
    return resultado, openpyxl.load_workbook(destino)


def test_cria_a_aba_do_topico_e_mantem_o_molde(gerado):
    resultado, workbook = gerado
    assert [p.sheet for p in resultado.synth] == ["REVESTIMENTO"]
    assert resultado.plans == []                     # nenhuma aba pré-existente
    assert "REVESTIMENTO" in workbook.sheetnames
    assert "MODELO BASE" in workbook.sheetnames      # molde intacto
    assert workbook["MODELO BASE"]["F7"].value is None


def test_cabecalho_de_tres_linhas_completo(gerado):
    """O molde real não tem fórmula de nome; a aba gerada nunca sai sem ela."""
    _, workbook = gerado
    sheet = workbook["REVESTIMENTO"]
    assert sheet["F7"].value == "X001"                       # código
    assert sheet["F8"].value == "CIMENTO"                    # nome (texto, D4)
    assert sheet["F9"].value.startswith("=IFERROR(VLOOKUP(F7,")   # unidade


def test_uma_linha_por_item_na_ordem_do_orcamento(gerado):
    _, workbook = gerado
    sheet = workbook["REVESTIMENTO"]
    assert [sheet[f"B{r}"].value for r in (10, 11, 12)] == ["S001", "S002", "S005"]
    assert [sheet[f"A{r}"].value for r in (10, 11, 12)] == ["5.1", "5.2", "5.3"]
    assert sheet["E10"].value == 120 and sheet["D10"].value == "M2"


def test_coeficiente_aninhado_vira_formula_rastreavel(gerado):
    _, workbook = gerado
    sheet = workbook["REVESTIMENTO"]
    assert sheet["F10"].value == "='COMPOSICOES'!D5"          # direto
    assert sheet["F11"].value == "='COMPOSICOES'!D16*0.025"   # via S003


def test_coluna_acrescentada_alem_do_molde(gerado):
    """O molde tem 2 lugares (F, G); o tópico precisa de 3."""
    resultado, workbook = gerado
    plano = resultado.synth[0]
    assert [letra for letra, _ in plano.columns] == ["F", "G", "H"]
    assert plano.appended == ["H"]
    assert workbook["REVESTIMENTO"]["H7"].value == "M001"


def test_coluna_acrescentada_herda_o_formato_da_coluna_modelo(template_workbook_path,
                                                              topics, tmp_path):
    import re
    import zipfile
    from orcauto.pipeline import run
    config = Config(compositions=CompositionConfig(sheets=["COMPOSICOES"],
                                                   service_code_re=r"^S"),
                    targets=TargetConfig(template_sheet="MODELO BASE"))
    destino = tmp_path / "estilo.xlsx"
    run(template_workbook_path, None, destino, config, topics=topics)
    with zipfile.ZipFile(destino) as pacote:
        alvo = [n for n in pacote.namelist() if "sheetAUTO" in n][0]
        xml = pacote.read(alvo).decode("utf-8")
    largura = re.search(r'<col min="8" max="8"[^>]*/>', xml)
    assert largura is not None                       # coluna H ganhou <col> próprio
    assert 'ref="A1:H' in re.search(r'<dimension ref="[^"]+"', xml).group(0)


def test_linha_de_total_desce_com_o_bloco_e_gera_a_formula(gerado):
    """O molde soma E10:E12 e o TOTAL fica na 14; com 3 itens ele sobe para a 13."""
    _, workbook = gerado
    sheet = workbook["REVESTIMENTO"]
    assert sheet["A13"].value == "TOTAL"
    assert sheet["F13"].value == "=SUMPRODUCT($E$10:$E$12,F10:F12)"
    assert sheet["A14"].value in (None, "")          # TOTAL antigo foi limpo


def test_insumo_ensacado_recebe_roundup_no_total(template_workbook_path, topics, tmp_path):
    from orcauto.pipeline import run
    config = Config(compositions=CompositionConfig(sheets=["COMPOSICOES"],
                                                   service_code_re=r"^S",
                                                   bag_rounding_inputs=("X001",)),
                    targets=TargetConfig(template_sheet="MODELO BASE"))
    destino = tmp_path / "saco.xlsx"
    run(template_workbook_path, None, destino, config, topics=topics)
    sheet = openpyxl.load_workbook(destino)["REVESTIMENTO"]
    assert sheet["F13"].value == "=ROUNDUP((SUMPRODUCT($E$10:$E$12,F10:F12)/50),0)"
    assert sheet["G13"].value == "=SUMPRODUCT($E$10:$E$12,G10:G12)"   # os demais, não


def test_item_sem_composicao_e_reportado_e_nao_vira_linha(gerado):
    resultado, _ = gerado
    plano = resultado.synth[0]
    assert [s.item.code for s in plano.skipped] == ["S999"]
    assert len(plano.rows) == 3


def test_geracao_pode_ser_desligada(template_workbook_path, topics, tmp_path):
    from orcauto.pipeline import run
    config = Config(compositions=CompositionConfig(sheets=["COMPOSICOES"],
                                                   service_code_re=r"^S"),
                    targets=TargetConfig(template_sheet="MODELO BASE",
                                         generate_missing=False))
    with pytest.raises(ValueError, match="nenhum tópico"):
        run(template_workbook_path, None, tmp_path / "x.xlsx", config, topics=topics)


def test_filtro_por_secao_deixa_so_materiais(resolver, topics):
    """Numa relação de materiais, mão de obra não deveria virar coluna."""
    index, resolve = resolver
    todas = {c.code for c in topic_inputs(topics[0].items, index, resolve)}
    materiais = {c.code for c in topic_inputs(topics[0].items, index, resolve,
                                              sections=("MATERIAIS",))}
    assert "M001" in todas and "M001" not in materiais      # PEDREIRO fica de fora
    assert {"X001", "X002"} <= materiais


def test_coluna_guarda_a_secao_de_origem(resolver, topics):
    index, resolve = resolver
    por_codigo = {c.code: c for c in topic_inputs(topics[0].items, index, resolve)}
    assert por_codigo["X001"].section == "MATERIAIS"
    assert por_codigo["M001"].section == "MAO DE OBRA"


def test_servico_sem_composicao_e_marcado(resolver, topics):
    """Um código que parece serviço mas não tem tabela vira coluna — sinalizada.

    No orçamento real é o C1603: o resolver não tem como abri-lo, então ele
    entra como folha. Silenciar isso esconderia um insumo que na verdade é um
    serviço inteiro não detalhado.
    """
    index, resolve = resolver
    por_codigo = {c.code: c for c in topic_inputs(topics[0].items, index, resolve)}
    assert por_codigo["X001"].unresolved_service is False


def test_topico_sem_coluna_nao_vira_aba_vazia(template_workbook_path, topics, tmp_path):
    """Com filtro de seção, um tópico só de mão de obra fica sem coluna alguma.

    Criar a aba assim seria ruído: ela registra serviços mas não levanta nada.
    O plano é mantido no relatório, com o motivo.
    """
    from orcauto.pipeline import run
    config = Config(compositions=CompositionConfig(sheets=["COMPOSICOES"],
                                                   service_code_re=r"^S",
                                                   column_sections=("EQUIPAMENTOS",)),
                    targets=TargetConfig(template_sheet="MODELO BASE"))
    destino = tmp_path / "vazia.xlsx"
    resultado = run(template_workbook_path, None, destino, config, topics=topics)
    assert [p.created for p in resultado.synth] == [False]
    assert "seções configuradas" in resultado.synth[0].reason
    assert "REVESTIMENTO" not in openpyxl.load_workbook(destino).sheetnames


# =========================================================================
# Fase 2.1 — a linha de TOTAL é reservada, e o formato é uniforme
# =========================================================================
def _muitos_itens():
    """Tópico maior que o bloco do molde (3 linhas), forçando transbordo."""
    from orcauto.pdf_budget import BudgetItem, Topic

    def item(ordem, codigo, quantidade):
        return BudgetItem(ordem, codigo, "X", "M2", quantidade, 1.0, quantidade,
                          1, 5, "REVESTIMENTO")

    return [Topic(5, "REVESTIMENTO", [
        item("5.1", "S001", 10.0), item("5.2", "S002", 20.0),
        item("5.3", "S005", 30.0), item("5.4", "S001", 40.0),
        item("5.5", "S002", 50.0),
    ])]


def _gerar(caminho, topicos, tmp_path, nome="saida.xlsx"):
    from orcauto.pipeline import run
    config = Config(compositions=CompositionConfig(sheets=["COMPOSICOES"],
                                                   service_code_re=r"^S"),
                    targets=TargetConfig(template_sheet="MODELO BASE"))
    destino = tmp_path / nome
    resultado = run(caminho, None, destino, config, topics=topicos)
    return resultado, openpyxl.load_workbook(destino)["REVESTIMENTO"]


def test_total_nunca_recebe_dados_de_item_ao_transbordar(template_workbook_path, tmp_path):
    """Com 5 itens e bloco de 3, um item cairia sobre a linha de TOTAL do molde.

    Era o defeito: a linha 14 continuava cinza, mas passava a exibir item,
    código e coeficientes, e o TOTAL de verdade ia parar mais abaixo sem
    formato nenhum.
    """
    _, sheet = _gerar(template_workbook_path, _muitos_itens(), tmp_path)
    assert [sheet[f"B{r}"].value for r in range(10, 15)] == \
           ["S001", "S002", "S005", "S001", "S002"]
    assert sheet["A15"].value == "TOTAL"
    assert sheet["B15"].value in (None, "")          # TOTAL não carrega código
    assert sheet["F15"].value == "=SUMPRODUCT($E$10:$E$14,F10:F14)"


def test_todas_as_linhas_de_item_tem_o_mesmo_formato(template_workbook_path, tmp_path):
    """Inclui as que caem sobre linhas que o molde já trazia com outro formato."""
    _, sheet = _gerar(template_workbook_path, _muitos_itens(), tmp_path)
    referencia = sheet["A10"]._style
    for row in range(11, 15):
        assert sheet[f"A{row}"]._style == referencia, f"linha {row} destoa"
        assert sheet[f"F{row}"]._style == sheet["F10"]._style


def test_total_deslocado_mantem_o_formato_de_total(template_workbook_path, tmp_path):
    _, sheet = _gerar(template_workbook_path, _muitos_itens(), tmp_path)
    assert sheet["A15"].font.bold is True
    assert sheet["A15"].fill.start_color.rgb == "FFD9D9D9"
    assert sheet["A15"]._style != sheet["A10"]._style      # não é linha de item


def test_sobra_do_molde_fica_oculta_quando_ha_poucos_itens(template_workbook_path,
                                                           topics, tmp_path):
    """3 itens num bloco de 3: o TOTAL sobe para a 13 e a 14 some da vista."""
    _, sheet = _gerar(template_workbook_path, topics, tmp_path)
    assert sheet["A13"].value == "TOTAL"
    assert sheet.row_dimensions[14].hidden is True
    assert sheet["A14"].value in (None, "")


def test_barra_do_cabecalho_recebe_o_nome_do_topico(template_workbook_path,
                                                    topics, tmp_path):
    """Vale tanto com transbordo quanto sem — não é efeito colateral do resto."""
    _, curto = _gerar(template_workbook_path, topics, tmp_path, "curto.xlsx")
    _, longo = _gerar(template_workbook_path, _muitos_itens(), tmp_path, "longo.xlsx")
    assert curto["A6"].value == "LEVANTAMENTO - REVESTIMENTO"
    assert longo["A6"].value == "LEVANTAMENTO - REVESTIMENTO"


def test_total_cobre_exatamente_as_linhas_de_item(template_workbook_path, tmp_path):
    resultado, sheet = _gerar(template_workbook_path, _muitos_itens(), tmp_path)
    plano = resultado.synth[0]
    assert (plano.first_row, plano.total_row) == (10, 15)
    assert sheet["G15"].value == "=ROUNDUP((SUMPRODUCT($E$10:$E$14,G10:G14)/50),0)" \
        or sheet["G15"].value == "=SUMPRODUCT($E$10:$E$14,G10:G14)"
