"""Integração de ponta a ponta sobre a pasta sintética (sem PDF e sem os arquivos reais)."""
import zipfile

import openpyxl
import pytest

from orcauto.config import CompositionConfig, Config, RulesConfig, TargetConfig
from orcauto.pipeline import match_topics, run, topic_sheet_score


@pytest.fixture
def config():
    return Config(compositions=CompositionConfig(sheets=["COMPOSICOES"], service_code_re=r"^S"),
                  targets=TargetConfig(), rules=RulesConfig())


@pytest.fixture
def generated(workbook_path, topics, config, tmp_path):
    destination = tmp_path / "saida.xlsx"
    result = run(workbook_path, None, destination, config, topics=topics)
    return result, openpyxl.load_workbook(destination)


def test_cria_a_aba_auto_e_mantem_a_original(generated):
    _, workbook = generated
    assert "REVESTIMENTO (AUTO)" in workbook.sheetnames
    assert "REVESTIMENTO" in workbook.sheetnames
    assert workbook["REVESTIMENTO"]["B11"].value == "S004"      # legado intacto


def test_grava_formula_rastreavel_e_nao_numero_solto(generated):
    _, workbook = generated
    sheet = workbook["REVESTIMENTO (AUTO)"]
    assert sheet["F10"].value == "='COMPOSICOES'!D5"
    assert sheet["F11"].value == "='COMPOSICOES'!D16*0.025"     # aninhado


def test_linha_atualizada_preserva_quantidade_da_aba(generated):
    _, workbook = generated
    assert workbook["REVESTIMENTO (AUTO)"]["E10"].value == 100


def test_linha_substituida_assume_o_item_do_orcamento(generated):
    _, workbook = generated
    sheet = workbook["REVESTIMENTO (AUTO)"]
    assert sheet["B11"].value == "S002"
    assert sheet["A11"].value == "5.2"
    assert sheet["E11"].value == 210


def test_linha_nova_vai_para_a_primeira_livre(generated):
    _, workbook = generated
    sheet = workbook["REVESTIMENTO (AUTO)"]
    assert sheet["B12"].value == "S005"
    assert sheet["E12"].value == 15


def test_total_intocado_quando_o_bloco_cabe(workbook_path, generated):
    _, workbook = generated
    original = openpyxl.load_workbook(workbook_path)
    assert workbook["REVESTIMENTO (AUTO)"]["F16"].value == original["REVESTIMENTO"]["F16"].value


def test_partes_originais_preservadas(workbook_path, generated):
    result, _ = generated
    original, output = zipfile.ZipFile(workbook_path), zipfile.ZipFile(result.output)
    permitidas = {"[Content_Types].xml", "xl/_rels/workbook.xml.rels",
                  "xl/workbook.xml", "xl/styles.xml"}
    alteradas = {n for n in set(original.namelist()) & set(output.namelist())
                 if original.read(n) != output.read(n)}
    assert alteradas <= permitidas


def test_aba_de_log_registra_o_mapeamento(generated):
    _, workbook = generated
    valores = [[c for c in row if c not in (None, "")]
               for row in workbook["LOG AUTO"].iter_rows(values_only=True)]
    texto = "\n".join(" | ".join(str(c) for c in linha) for linha in valores if linha)
    assert "1) MAPEAMENTO APLICADO" in texto
    assert "SUBSTITUI" in texto and "S004" in texto
    assert "S999" in texto                                   # item não aplicado


def test_amplia_o_total_quando_o_bloco_transborda(workbook_path, topics, config, tmp_path):
    """Com o bloco cheio, a linha nova cai fora do intervalo e o TOTAL acompanha."""
    workbook = openpyxl.load_workbook(workbook_path)
    sheet = workbook["REVESTIMENTO"]
    for row in (12, 13):                         # ocupa as duas linhas livres do intervalo
        sheet[f"B{row}"], sheet[f"E{row}"], sheet[f"F{row}"] = f"OUTRO{row}", 7, 1
    workbook.save(workbook_path)

    destination = tmp_path / "transborda.xlsx"
    result = run(workbook_path, None, destination, config, topics=topics)
    plan = result.plans[0]
    assert plan.new_last_row == 14               # S005 caiu fora do intervalo original
    gerado = openpyxl.load_workbook(destination)["REVESTIMENTO (AUTO)"]
    assert gerado["F16"].value == "=SUMPRODUCT($E$10:$E$14,F10:F14)"
    assert gerado["B14"].value == "S005"


def test_erro_util_quando_nenhum_topico_casa(workbook_path, topics, tmp_path):
    config = Config(compositions=CompositionConfig(sheets=["COMPOSICOES"]),
                    targets=TargetConfig(sheets=["NAO_EXISTE"]))
    with pytest.raises(ValueError, match="nenhum tópico"):
        run(workbook_path, None, tmp_path / "x.xlsx", config, topics=topics)


def test_pdf_ou_topics_e_obrigatorio(workbook_path, config, tmp_path):
    with pytest.raises(ValueError, match="pdf_path"):
        run(workbook_path, None, tmp_path / "x.xlsx", config)


def test_topic_sheet_score_nao_casa_por_substring():
    assert topic_sheet_score("SERVIÇOS PRELIMINARES", "RES") < 0.5
    assert topic_sheet_score("REVESTIMENTO", "REVESTIMENTOS") > 0.9
    assert topic_sheet_score("PAREDES E PAINÉIS", "PAREDES") > 0.9


def test_match_topics_respeita_o_mapa_explicito(topics):
    config = Config(targets=TargetConfig(topic_map={"5": "REVESTIMENTO"}))
    assert match_topics(topics, ["REVESTIMENTO", "OUTRA"], config) == [(topics[0], "REVESTIMENTO")]


def test_preenche_unidade_vazia_em_linha_atualizada(workbook_path, topics, config, tmp_path):
    """A planilha real tem VLOOKUP quebrado (#REF!) na coluna de unidade.

    Numa linha ATUALIZADA a automação não reescreve a identidade, então a
    unidade ficava vazia. Passa a preencher — só quando não há nada útil lá.
    """
    # linha 10 é ATUALIZADA (mesmo código do item 5.1)
    workbook = openpyxl.load_workbook(workbook_path)
    workbook["REVESTIMENTO"]["D10"] = None
    workbook.save(workbook_path)
    run(workbook_path, None, tmp_path / "vazia.xlsx", config, topics=topics)
    vazia = openpyxl.load_workbook(tmp_path / "vazia.xlsx")["REVESTIMENTO (AUTO)"]
    assert vazia["D10"].value == "M2"             # estava vazia: preenchida

    workbook = openpyxl.load_workbook(workbook_path)
    workbook["REVESTIMENTO"]["D10"] = "UN-MANUAL"
    workbook.save(workbook_path)
    run(workbook_path, None, tmp_path / "cheia.xlsx", config, topics=topics)
    cheia = openpyxl.load_workbook(tmp_path / "cheia.xlsx")["REVESTIMENTO (AUTO)"]
    assert cheia["D10"].value == "UN-MANUAL"      # tinha conteúdo: intocada


def test_erro_ref_conta_como_celula_vazia(workbook_path, topics, config, tmp_path):
    """`#REF!` não é conteúdo aproveitável — pode ser substituído."""
    from orcauto.layout import is_empty
    assert is_empty("#REF!") and is_empty("") and is_empty(None)
    assert not is_empty("M2") and not is_empty(0)


def test_log_lista_linhas_sem_contrapartida_no_orcamento(workbook_path, topics, config, tmp_path):
    """A linha legada que sobra precisa aparecer, não ficar invisível."""
    workbook = openpyxl.load_workbook(workbook_path)
    workbook["REVESTIMENTO"]["B12"] = "S777"
    workbook["REVESTIMENTO"]["C12"] = "SERVICO ANTIGO SEM RELACAO"
    workbook["REVESTIMENTO"]["E12"] = 50
    workbook.save(workbook_path)

    destination = tmp_path / "legado.xlsx"
    run(workbook_path, None, destination, config, topics=topics)
    linhas = ["|".join(str(c) for c in row if c not in (None, ""))
              for row in openpyxl.load_workbook(destination)["LOG AUTO"].iter_rows(values_only=True)]
    texto = "\n".join(linhas)
    assert "5) LINHAS DA ABA SEM CONTRAPARTIDA NO ORCAMENTO" in texto
    assert "S777" in texto
    assert "nao consta em nenhum item do orcamento" in texto


# ---------------------------------------------------------------------------
# Bug D do relatório v3 — o código do PDF aponta para outro serviço na tabela.
# O item pedido some da planilha e no lugar dele entra um que o PDF não pediu.
# Antes isso acontecia em silêncio: a descrição gravada vem da composição.
# ---------------------------------------------------------------------------

def test_descricao_divergente_entre_pdf_e_composicao_vira_aviso(
        workbook_path, config, tmp_path):
    """Caso real: o PDF diz que C0711 é CARGA MECANIZADA DE ENTULHO; a tabela
    do arquivo diz que C0711 é CARGA, DESCARGA E TRANSP. DE TUBOS DN 150mm."""
    from orcauto.pdf_budget import BudgetItem, Topic
    trocado = Topic(5, "REVESTIMENTO", [
        BudgetItem("5.1", "S001", "CARGA MECANIZADA DE ENTULHO EM CAMINHAO",
                   "M3", 120.0, 1.0, 120.0, 1, 5, "REVESTIMENTO"),
    ])
    destino = tmp_path / "divergente.xlsx"
    resultado = run(workbook_path, None, destino, config, topics=[trocado])

    divergencias = list(resultado.audit.divergences())
    assert len(divergencias) == 1
    _, item, composition, score = divergencias[0]
    assert item.code == "S001"
    assert composition.description == "CHAPISCO DE CIMENTO E AREIA"
    assert score < 0.60
    assert "descricao divergente" in resultado.report

    texto = "\n".join("|".join(str(c) for c in row if c not in (None, ""))
                      for row in openpyxl.load_workbook(destino)["LOG AUTO"]
                      .iter_rows(values_only=True))
    assert "7) DIVERGENCIA ENTRE A DESCRICAO DO PDF E A DA COMPOSICAO" in texto
    assert "CARGA MECANIZADA DE ENTULHO EM CAMINHAO" in texto


def test_descricao_equivalente_nao_vira_aviso(workbook_path, topics, config, tmp_path):
    resultado = run(workbook_path, None, tmp_path / "ok.xlsx", config, topics=topics)
    assert list(resultado.audit.divergences()) == []


def test_aviso_de_titulo_nao_reconhecido_chega_ao_log(workbook_path, topics, config, tmp_path):
    """A linha suspeita da aba de composições precisa aparecer na seção 8."""
    workbook = openpyxl.load_workbook(workbook_path)
    workbook["COMPOSICOES"]["A30"] = "RELATORIO ANALITICO - COMPOSICOES DE CUSTOS"
    workbook.save(workbook_path)

    destino = tmp_path / "aviso.xlsx"
    resultado = run(workbook_path, None, destino, config, topics=topics)
    assert any("A30" in aviso for aviso in resultado.audit.warnings)
    texto = "\n".join("|".join(str(c) for c in row if c not in (None, ""))
                      for row in openpyxl.load_workbook(destino)["LOG AUTO"]
                      .iter_rows(values_only=True))
    assert "8) AVISOS DE LEITURA DAS COMPOSICOES" in texto
    assert "RELATORIO ANALITICO" in texto
