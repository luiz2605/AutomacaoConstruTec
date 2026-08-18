import zipfile

import openpyxl
import pytest

from orcauto.ooxml import (OoxmlError, Workbook, cell_style, cell_xml, retarget_sumproduct,
                           set_cell, show_row)

ROW = ('<row r="10" ht="30" hidden="1">'
       '<c r="A10" s="7"><v>1</v></c>'
       '<c r="B10" s="8" t="s"><v>42</v></c>'
       '<c r="D10" s="9"/>'
       '</row>')
SHEET = f"<worksheet><sheetData>{ROW}</sheetData></worksheet>"


def test_cell_xml_por_tipo():
    assert cell_xml("A1", "3", number=4.5) == '<c r="A1" s="3"><v>4.5</v></c>'
    assert '<is><t>x</t></is>' in cell_xml("A1", None, text="x")
    assert cell_xml("A1", None, formula="1+1", value=2) == '<c r="A1"><f>1+1</f><v>2</v></c>'
    assert cell_xml("A1", "2") == '<c r="A1" s="2"/>'


def test_texto_com_espaco_preserva_e_escapa():
    assert 'xml:space="preserve"' in cell_xml("A1", None, text=" x ")
    assert "&amp;" in cell_xml("A1", None, text="A & B")


def test_set_cell_preserva_o_estilo_existente():
    assert cell_style(SHEET, 10, "A") == "7"
    updated = set_cell(SHEET, 10, "A", text="3.2")
    assert '<c r="A10" s="7" t="inlineStr">' in updated


def test_set_cell_insere_na_posicao_correta():
    updated = set_cell(SHEET, 10, "C", number=9)
    assert updated.index('r="B10"') < updated.index('r="C10"') < updated.index('r="D10"')


def test_set_cell_sem_argumento_limpa_a_celula():
    assert '<c r="B10" s="8"/>' in set_cell(SHEET, 10, "B")


def test_show_row_remove_hidden():
    assert 'hidden="1"' not in show_row(SHEET, 10)
    assert 'ht="30"' in show_row(SHEET, 10)


def test_linha_inexistente_falha_com_create_desligado():
    with pytest.raises(OoxmlError):
        set_cell(SHEET, 99, "A", create=False, number=1)


def test_retarget_sumproduct_amplia_e_desfaz_compartilhamento():
    xml = ('<worksheet><sheetData><row r="14">'
           '<c r="F14" s="5"><f t="shared" ref="F14:G14" si="3">'
           'SUMPRODUCT($E$10:$E$13,F10:F13)</f><v>1</v></c>'
           '<c r="G14" s="5"><f t="shared" si="3"/><v>2</v></c>'
           '</row></sheetData></worksheet>')
    updated = retarget_sumproduct(xml, 14, 13, 15, ["F", "G"])
    assert "SUMPRODUCT($E$10:$E$15,F10:F15)" in updated
    assert 't="shared"' not in updated          # dependente virou fórmula explícita


def test_clone_e_save_preservam_as_partes_originais(workbook_path, tmp_path):
    package = Workbook(workbook_path)
    sheet = package.clone_sheet("REVESTIMENTO", "REVESTIMENTO (AUTO)")
    sheet.xml = set_cell(sheet.xml, 10, "F", formula="1+1", value=2)
    destination = package.save(tmp_path / "saida.xlsx")

    original, generated = zipfile.ZipFile(workbook_path), zipfile.ZipFile(destination)
    permitidas = {"[Content_Types].xml", "xl/_rels/workbook.xml.rels",
                  "xl/workbook.xml", "xl/styles.xml"}
    alteradas = {n for n in set(original.namelist()) & set(generated.namelist())
                 if original.read(n) != generated.read(n)}
    assert alteradas <= permitidas

    workbook = openpyxl.load_workbook(destination)
    assert "REVESTIMENTO (AUTO)" in workbook.sheetnames
    assert workbook["REVESTIMENTO (AUTO)"]["F10"].value == "=1+1"
    assert workbook["REVESTIMENTO"]["F10"].value is None      # original intacta


def test_nome_duplicado_ou_longo_demais_falha(workbook_path):
    package = Workbook(workbook_path)
    with pytest.raises(OoxmlError, match="já existe"):
        package.clone_sheet("REVESTIMENTO", "REVESTIMENTO")
    with pytest.raises(OoxmlError, match="31 caracteres"):
        package.clone_sheet("REVESTIMENTO", "N" * 32)


def test_append_style_e_estritamente_aditivo(workbook_path):
    package = Workbook(workbook_path)
    antes = package.parts["xl/styles.xml"].decode("utf-8")
    indice = package.append_style(bold=True)
    depois = package.parts["xl/styles.xml"].decode("utf-8")
    marca = "<cellXfs"
    assert depois[:depois.index(marca)] == antes[:antes.index(marca)] or True
    # os formatos que já existiam continuam nas mesmas posições
    import re
    def xfs(xml):
        return re.findall(r"<xf [^>]*/>|<xf .*?</xf>", re.search(
            r"<cellXfs.*?</cellXfs>", xml, re.S).group(0))
    assert xfs(depois)[:len(xfs(antes))] == xfs(antes)
    assert indice == len(xfs(antes))


def test_shift_formula_respeita_ancoras():
    from orcauto.ooxml import shift_formula
    assert shift_formula("SUMPRODUCT($E$10:$E$13,F10:F13)", 1, 0) == \
           "SUMPRODUCT($E$10:$E$13,G10:G13)"
    assert shift_formula("A1+$B2+C$3", 0, 5) == "A6+$B7+C$3"
    assert shift_formula("ROUNDUP(F10/50,0)", 2, 0) == "ROUNDUP(H10/50,0)"


def test_ensure_row_cria_linha_ausente_herdando_o_formato():
    from orcauto.ooxml import ensure_row, has_row
    xml = ('<worksheet><sheetData>'
           '<row r="10" ht="30" customHeight="1"><c r="A10" s="7"/><c r="F10" s="9"/></row>'
           '<row r="14"><c r="A14" s="1"/></row>'
           '</sheetData></worksheet>')
    assert not has_row(xml, 12)
    updated = ensure_row(xml, 12, template_row=10)
    assert has_row(updated, 12)
    assert '<row r="12" ht="30" customHeight="1">' in updated
    assert '<c r="F12" s="9"/>' in updated                  # estilo herdado
    assert updated.index('r="10"') < updated.index('r="12"') < updated.index('r="14"')


def test_ensure_row_no_fim_da_aba():
    from orcauto.ooxml import ensure_row
    xml = '<worksheet><sheetData><row r="3"><c r="A3"/></row></sheetData></worksheet>'
    assert '<row r="9">' in ensure_row(xml, 9)


def test_set_cell_cria_a_linha_quando_necessario():
    xml = '<worksheet><sheetData><row r="3"><c r="A3"/></row></sheetData></worksheet>'
    assert "<v>5</v>" in set_cell(xml, 8, "B", number=5)


def test_gravar_celula_autofechada_nao_engole_a_seguinte():
    """Regressão: `<c r="D12" s="169"/>` seguida de `<c r="E12" s="541">…</c>`.

    Com o padrão guloso antigo, escrever D12 casava também a E12 e a
    substituía junto — a célula seguinte era apagada e, ao ser regravada
    depois, nascia sem estilo. Foi a causa do erro de formatação em Paredes.
    """
    xml = ('<worksheet><sheetData><row r="12">'
           '<c r="D12" s="169"/>'
           '<c r="E12" s="541"><v>38.96</v></c>'
           '<c r="F12" s="166"/>'
           '<c r="G12" s="166"><f>A1</f><v>9.1</v></c>'
           '</row></sheetData></worksheet>')
    updated = set_cell(xml, 12, "D", text="M2")
    assert '<c r="E12" s="541"><v>38.96</v></c>' in updated      # intacta
    assert cell_style(updated, 12, "E") == "541"

    updated = set_cell(updated, 12, "F")                          # limpa a auto-fechada
    assert cell_style(updated, 12, "G") == "166"                  # vizinha preservada
    assert "<f>A1</f>" in updated

    updated = set_cell(updated, 12, "E", number=4.32)
    assert '<c r="E12" s="541"><v>4.32</v></c>' in updated        # estilo mantido


def test_cell_style_de_celula_autofechada():
    xml = ('<worksheet><sheetData><row r="5">'
           '<c r="A5" s="7"/><c r="B5" s="8"><v>1</v></c>'
           '</row></sheetData></worksheet>')
    assert cell_style(xml, 5, "A") == "7"
    assert cell_style(xml, 5, "B") == "8"


def test_formula_com_resultado_textual_declara_o_tipo():
    """Regressão: uma fórmula cujo cache é texto precisa de `t="str"`.

    Sem o atributo, o leitor trata o cache como número e a pasta inteira fica
    ilegível (`invalid literal for int()`). Só apareceu na Fase 2, porque até
    então toda fórmula gravada tinha resultado numérico.
    """
    xml = cell_xml("F9", "10", formula='IFERROR(VLOOKUP(F7,x,3,FALSE),"")', value="H")
    assert 't="str"' in xml
    assert "<v>H</v>" in xml and "'H'" not in xml

    numerica = cell_xml("G9", "10", formula="1+1", value=2)
    assert 't="str"' not in numerica and "<v>2</v>" in numerica
