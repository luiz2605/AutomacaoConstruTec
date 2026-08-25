# -*- coding: utf-8 -*-
"""Fixtures sintéticas: uma pasta Excel e um orçamento criados do zero.

Nenhum teste depende dos arquivos reais do projeto — a suíte roda em qualquer
máquina, e o layout sintético é deliberadamente diferente do original (outras
linhas, outras colunas) para provar que nada está fixado em código.
"""
from __future__ import annotations

import sys
from pathlib import Path

import pytest

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "src"))

from orcauto.pdf_budget import BudgetItem, Topic          # noqa: E402

# Composições sintéticas: (linha, [valores da coluna A em diante])
COMPOSITION_ROWS = {
    1:  ["Relatorio de Composicoes"],
    3:  ["S001 - CHAPISCO DE CIMENTO E AREIA - M2"],
    4:  ["MATERIAIS", None, "Unidade", "Coeficiente"],
    5:  ["X001", "CIMENTO", "KG", 2.5],
    6:  ["X002", "AREIA", "M3", 0.006],
    8:  ["S002 - REBOCO COM ARGAMASSA TRACO 1:5 - M2"],
    9:  ["MAO DE OBRA", None, "Unidade", "Coeficiente"],
    10: ["M001", "PEDREIRO", "H", 0.6],
    11: ["SERVICOS"],
    12: ["S003", "ARGAMASSA DE CIMENTO E AREIA", "M3", 0.025],
    14: ["S003 - ARGAMASSA DE CIMENTO E AREIA - M3"],
    15: ["MATERIAIS", None, "Unidade", "Coeficiente"],
    16: ["X001", "CIMENTO", "KG", 292],
    17: ["X002", "AREIA", "M3", 1.216],
    # legado: mesmo serviço do S002, com traço diferente e fora do orçamento
    19: ["S004 - REBOCO COM ARGAMASSA TRACO 1:6 - M2"],
    20: ["SERVICOS", None, "Unidade", "Coeficiente"],
    21: ["S003", "ARGAMASSA DE CIMENTO E AREIA", "M3", 0.02],
    23: ["S005 - CONCRETO ESTRUTURAL - M3"],
    24: ["MATERIAIS", None, "Unidade", "Coeficiente"],
    25: ["X001", "CIMENTO", "KG", 350],
    26: ["X002", "AREIA", "M3", 0.87],
    27: ["SERVICOS"],
    # armadilha: S001 aparece aqui como SUB-ITEM, com coeficiente diferente do
    # da sua própria tabela. Pegar este valor seria o falso positivo clássico.
    28: ["S001", "CHAPISCO DE CIMENTO E AREIA", "M2", 3],
}


@pytest.fixture
def composition_rows() -> list[list]:
    """As mesmas composições como lista densa de linhas (para testes puros)."""
    last = max(COMPOSITION_ROWS)
    return [COMPOSITION_ROWS.get(number, []) for number in range(1, last + 1)]


@pytest.fixture
def workbook_path(tmp_path: Path) -> Path:
    """Pasta .xlsx sintética com uma aba de composições e uma de levantamento."""
    import openpyxl

    workbook = openpyxl.Workbook()
    compositions = workbook.active
    compositions.title = "COMPOSICOES"
    for number, values in COMPOSITION_ROWS.items():
        for offset, value in enumerate(values, start=1):
            if value is not None:
                compositions.cell(number, offset, value)

    target = workbook.create_sheet("REVESTIMENTO")
    target["A1"] = "RELACAO DE MATERIAIS"
    target["A6"] = "LEVANTAMENTO - REVESTIMENTO"
    target["F7"], target["G7"] = "X001", "X002"
    for column, header in zip("ABCDEFG",
                              ["ITEM", "CODIGO", "SERVICOS", "UN", "QUANT.", "CIMENTO", "AREIA"]):
        target[f"{column}8"] = header
    target["F9"], target["G9"] = "KG", "M3"
    target["A10"], target["B10"] = 1, "S001"
    target["C10"], target["D10"], target["E10"] = "CHAPISCO DE CIMENTO E AREIA", "M2", 100
    target["A11"], target["B11"] = 2, "S004"
    target["C11"], target["D11"], target["E11"] = "REBOCO COM ARGAMASSA TRACO 1:6", "M2", 200
    target["F11"], target["G11"] = 5.84, 0.0243
    # TOTAL propositalmente afastado do fim do intervalo somado: sobram as linhas
    # 14 e 15, utilizáveis apenas se o intervalo do TOTAL for ampliado.
    target["A16"] = "TOTAL"
    target["F16"] = "=SUMPRODUCT($E$10:$E$13,F10:F13)"
    target["G16"] = "=SUMPRODUCT($E$10:$E$13,G10:G13)"

    path = tmp_path / "levantamento.xlsx"
    workbook.save(path)
    return path


@pytest.fixture
def topics() -> list[Topic]:
    """Orçamento sintético equivalente ao tópico 5 de um orçamento real."""
    def item(order, code, description, unit, quantity):
        return BudgetItem(order, code, description, unit, quantity, 1.0, quantity,
                          1, 5, "REVESTIMENTO")

    topic = Topic(5, "REVESTIMENTO", [
        item("5.1", "S001", "CHAPISCO DE CIMENTO E AREIA", "M2", 120.0),
        item("5.2", "S002", "REBOCO COM ARGAMASSA TRACO 1:5", "M2", 210.0),
        item("5.3", "S005", "CONCRETO ESTRUTURAL", "M3", 15.0),
        item("5.4", "S999", "SERVICO SEM COMPOSICAO", "M2", 5.0),
    ])
    return [topic]


@pytest.fixture
def template_workbook_path(tmp_path: Path) -> Path:
    """Pasta com as composições e APENAS a aba-molde — nenhuma aba de destino.

    Reproduz a situação do arquivo-base real: `MODELO BASE` traz o desenho da
    tabela (cabeçalho de três linhas, bloco de serviços, linha de TOTAL) mas
    nenhum insumo declarado, e por isso não passa em `layout.detect`.
    """
    import openpyxl

    workbook = openpyxl.Workbook()
    compositions = workbook.active
    compositions.title = "COMPOSICOES"
    for number, values in COMPOSITION_ROWS.items():
        for offset, value in enumerate(values, start=1):
            if value is not None:
                compositions.cell(number, offset, value)

    template = workbook.create_sheet("MODELO BASE")
    template["A1"] = "RELACAO DE MATERIAIS"
    for column, header in zip("ABCDE", ["ITEM", "CODIGO", "SERVICOS", "UN", "QUANT."]):
        template[f"{column}8"] = header
    # duas colunas de insumo já desenhadas; o resto terá de ser acrescentado
    for column in ("F", "G"):
        template[f"{column}9"] = f'=IFERROR(VLOOKUP({column}7,insumos!$A$6:$D$99,3,FALSE),"")'
    template["A6"] = ""                       # barra colorida, vazia no molde
    template["A14"] = "TOTAL"
    template["F14"] = "=SUMPRODUCT($E$10:$E$12,F10:F12)"
    template["G14"] = "=ROUNDUP((SUMPRODUCT($E$10:$E$12,G10:G12)/50),0)"

    # formatos distintos, para que os testes consigam distinguir uma linha de
    # item de uma linha de TOTAL — é o que o molde real faz com o cinza
    from openpyxl.styles import Border, Font, PatternFill, Side
    borda = Border(*[Side(style="thin")] * 4)
    cinza = PatternFill("solid", start_color="FFD9D9D9")
    for column in "ABCDEFG":
        template[f"{column}10"].border = borda
        template[f"{column}14"].fill = cinza
        template[f"{column}14"].font = Font(bold=True)
    # a linha 13 do molde imita as linhas "de passagem" quebradas do real:
    # existe, mas com formato diferente das linhas de item
    template["A13"] = "=B99"

    # H e I existem na linha de item do molde, mas são de FORA da tabela — no
    # molde real a linha 10 vai até S com estilos residuais. Uma coluna
    # acrescentada que caia nessas posições não pode herdar isso.
    from openpyxl.styles import Alignment
    for column in ("H", "I"):
        template[f"{column}10"].alignment = Alignment(horizontal="center")
        template[f"{column}10"].number_format = "0.00"
        template[f"{column}14"].number_format = "General"

    path = tmp_path / "base.xlsx"
    workbook.save(path)
    return path
