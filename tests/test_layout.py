import openpyxl
import pytest

from orcauto.layout import LayoutError, detect


@pytest.fixture
def sheets(workbook_path):
    formulas = openpyxl.load_workbook(workbook_path)
    values = openpyxl.load_workbook(workbook_path, data_only=True)
    return formulas, values


def test_detecta_layout_sem_posicoes_fixas(sheets):
    formulas, values = sheets
    layout = detect(formulas["REVESTIMENTO"], values["REVESTIMENTO"])
    assert layout.header_row == 8
    assert layout.input_row == 7
    assert layout.quantity_column == "E"
    assert (layout.first_row, layout.last_row) == (10, 13)
    assert layout.total_row == 16
    assert layout.tracked == {"F": "X001", "G": "X002"}


def test_le_as_linhas_de_servico(sheets):
    formulas, values = sheets
    layout = detect(formulas["REVESTIMENTO"], values["REVESTIMENTO"])
    assert layout.rows[10].code == "S001"
    assert layout.rows[10].quantity == 100
    assert layout.rows[11].filled_columns == {"F", "G"}
    assert layout.free_rows() == [12, 13, 14, 15]
    assert layout.rows_beyond_total() == [14, 15]


def test_linha_com_escopo_manual_nao_e_livre(workbook_path):
    """Uma linha sem código, mas com quantidade, carrega escopo e não pode ser
    sobrescrita pela automação."""
    workbook = openpyxl.load_workbook(workbook_path)
    workbook["REVESTIMENTO"]["E12"] = 42
    workbook["REVESTIMENTO"]["C12"] = "INSUMO AVULSO"
    workbook.save(workbook_path)
    formulas = openpyxl.load_workbook(workbook_path)
    values = openpyxl.load_workbook(workbook_path, data_only=True)
    layout = detect(formulas["REVESTIMENTO"], values["REVESTIMENTO"])
    assert layout.free_rows() == [13, 14, 15]


def test_erro_quando_nao_ha_total(tmp_path):
    workbook = openpyxl.Workbook()
    workbook.active["A1"] = "sem levantamento aqui"
    path = tmp_path / "vazio.xlsx"
    workbook.save(path)
    formulas = openpyxl.load_workbook(path)
    values = openpyxl.load_workbook(path, data_only=True)
    with pytest.raises(LayoutError, match="TOTAL"):
        detect(formulas.active, values.active)


def test_erro_quando_o_total_nao_tem_sumproduct(tmp_path):
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet["A8"], sheet["B8"], sheet["E8"] = "ITEM", "CODIGO", "QUANT."
    sheet["A16"] = "TOTAL"
    sheet["F16"] = 123
    path = tmp_path / "sem_sumproduct.xlsx"
    workbook.save(path)
    formulas = openpyxl.load_workbook(path)
    values = openpyxl.load_workbook(path, data_only=True)
    with pytest.raises(LayoutError, match="SUMPRODUCT"):
        detect(formulas.active, values.active)
